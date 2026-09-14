# -*- coding: utf-8 -*-
"""
Разбор технологических регламентов на фрагменты для векторного поиска.

Чанк равен **нумерованному пункту** регламента, а не абзацу и не окну из N
символов. Так сделано потому, что пункт — это и есть единица нормы: на него
ссылаются, его цитируют, им доказывают. Координата фрагмента («ТР-ЭКС-2026/01
п. 4.3») попадает прямо в ответ агента, и её можно проверить, открыв документ.

Из раздела «Приложения» каждого регламента вынимается таблица соответствия
«номер приложения → таблица НСИ». Благодаря ей ссылка в тексте пункта
(«согласно приложению 5») превращается в номер таблицы НСИ (табл. 12) — это
основа графа связей `memory/links.py`.

Пункты про охрану труда, уборку помещений и заключительные положения помечаются
как не относящиеся к планированию: в выдачу поиска они по умолчанию не идут,
иначе вопрос про выбор линии вытягивает инструктаж по спецодежде.
"""
from __future__ import annotations

import re
from pathlib import Path

CLAUSE_RE = re.compile(r"^(\d+\.\d+)[.\s \t]+(\S.*)$")
SECTION_RE = re.compile(r"^(\d+)\.\s+(\S.*)$")
DOCNUM_RE = re.compile(r"Номер документа:\s*([^.]+)\.")
STAGE_RE = re.compile(r"Этап\s+(\w+)")
APPENDIX_REF_RE = re.compile(r"прилож\w*\s*((?:\d+\s*(?:,|и|\s)\s*)*\d+)", re.IGNORECASE)
TABLE_REF_RE = re.compile(r"табл\w*\.?\s*(\d+)", re.IGNORECASE)
TABLE_IN_CELL_RE = re.compile(r"(\d+)")

STAGE_NORM = {"экструзии": "экструзия", "печати": "печать", "кольцевания": "кольцевание"}
OFF_TOPIC = ("охран", "уборк", "обслуживание помещений", "заключительн", "общие положения")


def _appendix_rows(doc) -> list[dict]:
    """Строки раздела «Приложения»: номер приложения, наименование, таблица НСИ.

    Наименование берётся отсюда не случайно. Человеческое название таблицы НСИ
    («Печатное оборудование (допустимость линий)») не лежит ни в конфиге стенда,
    ни в самих файлах НСИ — там только ключ, имя файла и заголовки колонок. Оно
    есть ровно в одном месте: в колонке «Наименование приложения» регламента.
    Раньше эти названия были переписаны в словарь внутри инструмента — то есть
    агент показывал пользователю мои формулировки как вокабуляр стенда, и на
    другом стенде показывал бы их же.
    """
    rows: list[dict] = []
    for table in doc.tables:
        header = [c.text.strip().lower() for c in table.rows[0].cells]
        if not any("приложени" in h for h in header):
            continue
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            if len(cells) < 3:
                continue
            number = cells[0].strip()
            m = TABLE_IN_CELL_RE.search(cells[2])
            if number.isdigit() and m:
                rows.append({"приложение": number, "наименование": cells[1].strip(),
                             "таблица": m.group(1)})
    return rows


def _appendix_map(doc) -> dict[str, str]:
    """Приложение № → номер таблицы НСИ, из раздела «Приложения»."""
    return {r["приложение"]: r["таблица"] for r in _appendix_rows(doc)}


def _referenced_tables(text: str, appendix: dict[str, str]) -> list[str]:
    tables: list[str] = []
    for match in APPENDIX_REF_RE.finditer(text):
        for num in re.findall(r"\d+", match.group(1)):
            if num in appendix:
                tables.append(appendix[num])
    tables.extend(TABLE_REF_RE.findall(text))
    return sorted(set(tables), key=int)


def parse_regulation(path: Path) -> list[dict]:
    """Возвращает список фрагментов одного регламента."""
    import docx  # локальный импорт: python-docx нужен только на сборке индекса

    doc = docx.Document(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    head = "\n".join(paragraphs[:12])

    m = DOCNUM_RE.search(head)
    doc_number = m.group(1).strip() if m else path.stem
    doc_code = doc_number.split("-20")[0] if "-20" in doc_number else doc_number
    doc_title = next((p for p in paragraphs if p.startswith("Технологический регламент")),
                     path.stem)
    sm = STAGE_RE.search(doc_title)
    stage = STAGE_NORM.get((sm.group(1) if sm else "").lower(), "")
    appendix = _appendix_map(doc)

    chunks: list[dict] = []
    section_no, section_title = "0", "Без раздела"
    for text in paragraphs:
        clause = CLAUSE_RE.match(text)
        if clause:
            number, body = clause.group(1), clause.group(2).strip()
            relevant = not any(w in section_title.lower() for w in OFF_TOPIC[:-1])
            chunks.append({
                "id": f"{doc_code} п. {number}",
                "text": body,
                "meta": {
                    "источник": "regulations",
                    "документ": doc_code,
                    "номер_документа": doc_number,
                    "название": doc_title,
                    "этап": stage,
                    "раздел": f"{section_no}. {section_title}",
                    "пункт": number,
                    "таблицы_НСИ": _referenced_tables(body, appendix),
                    "координата": f"{doc_number} п. {number}",
                    "к_планированию": relevant,
                    "файл": path.name,
                },
            })
            continue
        section = SECTION_RE.match(text)
        if section:
            section_no, section_title = section.group(1), section.group(2).strip()

    return chunks


def embedding_text(chunk: dict) -> str:
    """Текст, который уходит в эмбеддер: пункт вместе с контекстом раздела."""
    m = chunk["meta"]
    tables = (" Ссылается на таблицы НСИ: " + ", ".join(m["таблицы_НСИ"]) + "."
              if m["таблицы_НСИ"] else "")
    return (f"{m['название']}. Этап: {m['этап']}. Раздел {m['раздел']}. "
            f"Пункт {m['пункт']}: {chunk['text']}{tables}")


def collect(cfg) -> tuple[list[dict], dict[str, dict[str, str]]]:
    """Все фрагменты регламентов стенда и карты приложений по документам."""
    import docx

    chunks: list[dict] = []
    appendices: dict[str, dict[str, str]] = {}
    for path in sorted(cfg.stand.reglaments_dir.glob("*.docx")):
        if path.name.startswith("~$"):      # локи Word — не документы
            continue
        part = parse_regulation(path)
        chunks.extend(part)
        if part:
            appendices[part[0]["meta"]["документ"]] = _appendix_map(docx.Document(str(path)))
    return chunks, appendices


def table_titles(cfg) -> dict[str, dict]:
    """Номер таблицы НСИ → как её называет регламент, и где он это делает."""
    import docx  # локальный импорт: python-docx нужен только на сборке индекса

    out: dict[str, dict] = {}
    for path in sorted(cfg.stand.reglaments_dir.glob("*.docx")):
        if path.name.startswith("~$"):
            continue
        doc = docx.Document(str(path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        m = DOCNUM_RE.search("\n".join(paragraphs[:12]))
        number_doc = m.group(1).strip() if m else path.stem
        for row in _appendix_rows(doc):
            if not row["наименование"]:
                continue
            out.setdefault(row["таблица"], {
                "название": row["наименование"],
                "приложение": row["приложение"],
                "документ": number_doc,
                "координата": f"{number_doc}, приложение {row['приложение']}",
            })
    return out
